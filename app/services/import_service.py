# -*- coding: utf-8 -*-
"""import_service —— Excel 解析流水线（M2.1 实现，FastAPI/SQLAlchemy 版）。

职责（M1 §8）：只读打开 Excel → 解析（六步流水线 §8.1）→ 预览（不写库）
→ 确认入库（调 upsert_service，时序见 §8.6）。

设计来源：原 Odoo 版 services/import_service.py（算法 100% 继承，仅框架切换）。

解析流程：
1. 打开工作簿（read_only + data_only），可选指定 sheet；
2. 合并单元格填充（read_only 模式 merged_cells 可能为空，则跳过）；
3. 表头探测：逐行 map_headers，命中「项目名称」即表头行；
4. 必填列校验（item_name 缺失 → 整表终止，不静默导入垃圾）；
5. 逐数据行：别名映射 → 国标编码解析 → 单位归一化 → 数值字段；
6. 合计行检测（item_name ∈ 合计/总计/小计）→ 记为 summary_total，不入库；
7. 返回 ParsedResult 信封：rows / row_count / skipped_count / anomaly_count /
   warnings / parsed_total / summary_total / file_hash / sheet_name。

安全约束（S1 源文件只读）：
- load_workbook 必须 read_only=True, data_only=True；
- 禁止对源文件调用 wb.save()；
- 解析过程不修改源文件。
"""
import hashlib
import io
import logging
from typing import Optional, Dict, Any, List

import openpyxl

from data.aliases import map_headers, validate_required, extract_field_value
from data.gb_code import parse_gb_code
from data.unit_normalize import normalize_unit

logger = logging.getLogger("zaojia.import")

# 合计行关键词（命中的行不入库，其合价作为 summary_total 参与校验和比对）
SUMMARY_KEYWORDS = ('合计', '总计', '小计', '求和', '总结')

# 数值字段 → 对应 _num 字段
NUMERIC_FIELDS = [
    ('quantity', 'quantity_num'),
    ('unit_rate', 'unit_rate_num'),
    ('total', 'total_num'),
    ('provisional_sum', 'provisional_sum_num'),
]

# 其余文本字段
TEXT_FIELDS = ('project_name', 'sub_division', 'ordinal', 'source_path')


def parse_excel(
    file_bytes: Optional[bytes] = None,
    file_path: Optional[str] = None,
    filename: Optional[str] = None,
    sheet: Optional[str] = None,
) -> Dict[str, Any]:
    """解析 Excel → 结构化 ParsedResult 信封（不写库）。

    参数
    ----
    file_bytes : bytes, optional —— 上传文件二进制
    file_path : str, optional —— 本地文件路径
    filename : str, optional —— 来源文件名（用于批次名）
    sheet : str, optional —— 指定工作表名（默认取首个）

    返回
    ----
    dict : ParsedResult 信封
        {
            'file_hash': str, 'filename': str, 'sheet_name': str,
            'rows': list[dict], 'row_count': int,
            'skipped_count': int, 'anomaly_count': int,
            'warnings': list, 'parsed_total': float,
            'summary_total': float | None,
            'unrecognized_columns': list[str],
        }

    异常
    ----
    ValueError —— 未提供文件、工作表不存在、未找到表头、缺少必填列
    """
    if file_bytes is None and file_path is None:
        raise ValueError("未提供文件（file_bytes 或 file_path 至少其一）。")

    # 只读 + 只值打开（S1 源文件只读）
    if file_bytes is not None:
        stream = io.BytesIO(file_bytes)
        wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
    else:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

    try:
        file_hash = _sha256(file_bytes, file_path)

        # sheet 选择
        if sheet:
            if sheet not in wb.sheetnames:
                raise ValueError(f"工作表不存在: {sheet}")
            ws = wb[sheet]
        else:
            ws = wb[wb.sheetnames[0]]
            sheet = ws.title

        # 合并单元格填充表
        merge_map = _merged_map(ws)

        # 表头探测
        header_row_idx, header_map, multi_hit, unrecognized = _detect_header(ws, merge_map)
        if header_row_idx is None:
            raise ValueError("未找到表头行（必须包含「项目名称」列）。")
        missing = validate_required(header_map)
        if missing:
            raise ValueError(f"缺少必填列：{', '.join(missing)}")

        # 逐行解析
        rows = []
        warnings = []
        skipped = 0
        anomaly = 0
        parsed_total = 0.0
        summary_total = None
        max_row = ws.max_row or 0

        for r in range(header_row_idx + 1, max_row + 1):
            vals = _read_row(ws, r, header_map, multi_hit, merge_map)
            item_name = (vals.get('item_name') or '').strip()

            # 合计行：不入库，记录 summary_total
            if item_name and item_name in SUMMARY_KEYWORDS:
                tv = _to_float(vals.get('total'))
                if tv is not None:
                    summary_total = tv
                continue

            if not item_name:
                skipped += 1
                continue

            rec = {'sequence': r, 'source_sheet': sheet, 'item_name': item_name}
            row_warns = []

            # 国标编码（保留原始值 + 规范化完整值）
            raw_code = vals.get('item_code')
            if raw_code is not None and str(raw_code).strip():
                code_raw = str(raw_code).strip()
                rec['item_code_raw'] = code_raw
                _code_9, code_full = parse_gb_code(code_raw)
                rec['item_code'] = code_full

            # 项目特征描述
            feat = vals.get('item_feature')
            if feat is not None and str(feat).strip():
                rec['item_feature'] = str(feat).strip()

            # 计量单位 + 归一化
            unit = vals.get('unit')
            if unit is not None and str(unit).strip():
                rec['unit'] = str(unit).strip()
                unit_std, warn = normalize_unit(unit)
                rec['unit_std'] = unit_std
                if warn:
                    row_warns.append(f'计量单位未归一化: {unit}（建议补映射或人工确认）')

            # 数值字段
            for canon, numf in NUMERIC_FIELDS:
                v = vals.get(canon)
                if v is None or str(v).strip() == '':
                    continue
                rec[canon] = str(v).strip()
                fv = _to_float(v)
                if fv is not None:
                    rec[numf] = fv
                    if canon == 'total':
                        parsed_total += fv
                else:
                    row_warns.append(f'{canon} 非数值: {v}')

            # 其余文本字段
            for canon in TEXT_FIELDS:
                v = vals.get(canon)
                if v is not None and str(v).strip():
                    rec[canon] = str(v).strip()

            # 章节/分部标题行：有名称但无编码且无量价 → 结构性行，跳过不入库
            if (not rec.get('item_code')
                    and rec.get('quantity_num') is None
                    and rec.get('unit_rate_num') is None
                    and rec.get('total_num') is None):
                continue

            # 异常标记
            if row_warns:
                anomaly += 1
                rec['anomaly_flag'] = 'warning'
                rec['anomaly_reason'] = '；'.join(row_warns)[:500]
            else:
                rec['anomaly_flag'] = 'normal'

            for w in row_warns:
                warnings.append({
                    'row': r, 'reason': w,
                    'raw': rec.get('item_code_raw', ''),
                    'suggestion': '导入后人工核对',
                })
            rows.append(rec)

        return {
            'file_hash': file_hash,
            'filename': filename,
            'sheet_name': sheet,
            'rows': rows,
            'row_count': len(rows),
            'skipped_count': skipped,
            'anomaly_count': anomaly,
            'warnings': warnings,
            'parsed_total': round(parsed_total, 2),
            'summary_total': summary_total,
            'unrecognized_columns': [c for _, c in unrecognized],
        }
    finally:
        wb.close()


def preview_parsed(parsed: Dict[str, Any]) -> Dict[str, Any]:
    """解析结果预览信封（不写库，精简字段）。"""
    return {
        'row_count': parsed.get('row_count', 0),
        'skipped_count': parsed.get('skipped_count', 0),
        'anomaly_count': parsed.get('anomaly_count', 0),
        'warnings': parsed.get('warnings', []),
        'parsed_total': parsed.get('parsed_total', 0),
        'summary_total': parsed.get('summary_total'),
        'unrecognized_columns': parsed.get('unrecognized_columns', []),
    }


# ------------------------------------------------------------------
# 内部辅助
# ------------------------------------------------------------------

def _read_row(ws, r: int, header_map: Dict[int, str],
              multi_hit: Dict[str, List[int]], merge_map: Dict[tuple, Any]) -> Dict[str, Any]:
    """读取第 r 行，按 header_map 映射到规范字段（含合并单元格填充 + 多列命中回退）。

    header_map：键为 1 基列号（与 enumerate(start=1) 对齐）。
    multi_hit：键为规范字段名，值为 0 基列索引列表（extract_field_value 用 0 基）。
    """
    # 读取全部列 → 0 基列表
    row_values = []
    for col, cell in enumerate(next(ws.iter_rows(min_row=r, max_row=r)), start=1):
        val = cell.value
        if (val is None or val == '') and (r, col) in merge_map:
            val = merge_map[(r, col)]
        row_values.append(val)

    cells = {}
    for col_1based, canon in header_map.items():
        idx = col_1based - 1  # 1 基 → 0 基
        if canon in multi_hit:
            # 多列命中：用 extract_field_value 取首个非空列（0 基索引）
            cells[canon] = extract_field_value(row_values, multi_hit[canon])
        elif idx < len(row_values):
            cells[canon] = row_values[idx]
    return cells


def _read_full_row(ws, r: int, merge_map: Dict[tuple, Any]) -> Dict[int, Any]:
    """读取第 r 行全部列 → {col_index(1基): value}（含合并填充）。"""
    out = {}
    for col, cell in enumerate(next(ws.iter_rows(min_row=r, max_row=r)), start=1):
        val = cell.value
        if (val is None or val == '') and (r, col) in merge_map:
            val = merge_map[(r, col)]
        out[col] = val
    return out


def _detect_header(ws, merge_map: Dict[tuple, Any], max_scan: int = 15):
    """扫描前 max_scan 行，找到含 item_name 的表头行。

    返回 (header_row_idx, header_map(1基), multi_hit(0基), unrecognized)。
    header_map 的键为 1 基列号（与 _read_row 的 enumerate start=1 对齐）。
    multi_hit 的值为 0 基列索引（extract_field_value 用 0 基）。
    """
    max_row = min(max_scan, ws.max_row or 0)
    for r in range(1, max_row + 1):
        full = _read_full_row(ws, r, merge_map)
        row_values = [full.get(c) for c in sorted(full.keys())]
        header_map, unrecognized, multi_hit = map_headers(row_values)
        if 'item_name' in header_map.values():
            # header_map：0 基 → 1 基对齐（_read_row 用 enumerate start=1）
            header_map = {k + 1: v for k, v in header_map.items()}
            # multi_hit：保持 0 基（extract_field_value 用 0 基索引）
            return r, header_map, multi_hit, unrecognized
    return None, None, {}, []


def _merged_map(ws) -> Dict[tuple, Any]:
    """构建合并单元格 → 左上值 的映射（read_only 无 merged_cells 则返回空）。"""
    m = {}
    try:
        ranges = ws.merged_cells.ranges
    except Exception:
        return m
    try:
        for rng in ranges:
            tl = ws.cell(row=rng.min_row, column=rng.min_col).value
            for rr in range(rng.min_row, rng.max_row + 1):
                for cc in range(rng.min_col, rng.max_col + 1):
                    m[(rr, cc)] = tl
    except Exception:
        return m
    return m


def _to_float(v) -> Optional[float]:
    """单元格 → float（容忍千分位/货币符号/空值）。"""
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    s = str(v).strip().replace(',', '').replace('￥', '').replace('¥', '')
    if s == '':
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _sha256(file_bytes: Optional[bytes] = None, file_path: Optional[str] = None) -> str:
    """计算文件 SHA256（十六进制小写），用于批次幂等键。"""
    h = hashlib.sha256()
    if file_bytes is not None:
        h.update(file_bytes)
    else:
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(8192), b''):
                h.update(chunk)
    return h.hexdigest()
