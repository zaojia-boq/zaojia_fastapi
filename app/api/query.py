# -*- coding: utf-8 -*-
"""查询/导出 API（M2-查询导出 实现）。

路由：
- GET /api/query/items —— 分页查询清单项（默认 active=True）
- GET /api/query/items/{biz_id} —— 单条详情
- GET /api/query/export —— 导出 Excel（>1000 行需 confirm=true 二次确认）

权限：查询需登录（三角色均可）；导出限 estimator/admin（viewer 禁导出，M1.5 角色定义）；导出写审计日志（导出留痕，M2 §6.2）。
"""
import io
import json
import uuid
from datetime import datetime, timezone
from typing import Optional, List

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.db import get_db
from app.core.security import get_current_user, require_role, ROLE_ESTIMATOR, ROLE_ADMIN
from app.core.audit import log_audit, ACTION_EXPORT
from app.services.query_service import search_items, get_item, export_items

router = APIRouter(prefix="/api/query", tags=["查询导出"])

# 大批量导出闸门（M2 §6.2）：单次 > 1000 行需二次确认
EXPORT_GATE_LIMIT = 1000

# 导出字段（默认列，M2 §6.1）
EXPORT_COLUMNS = [
    ('item_code', '项目编码'),
    ('item_name', '项目名称'),
    ('item_feature', '项目特征描述'),
    ('unit', '计量单位'),
    ('quantity', '工程量'),
    ('unit_rate', '综合单价'),
    ('total', '合价'),
    ('std_name', '标准名称'),
    ('std_spec', '标准规格'),
    ('material_dict_id', '物料字典ID'),
    ('data_source_type', '数据性质'),
    ('province', '省份'),
    ('price_period', '价格期'),
    ('source_sheet', '来源Sheet'),
    ('source_path', '来源文件'),
]


def _parse_domain(filters: Optional[str]) -> List:
    """解析前端 filters 字符串 → domain 列表。

    支持格式：JSON 数组，如
        [["item_name","like","土方"],["data_source_type","=","completed"]]
    """
    if not filters:
        return []
    try:
        parsed = json.loads(filters)
        if not isinstance(parsed, list):
            return []
        return [c for c in parsed if isinstance(c, list) and len(c) >= 2]
    except json.JSONDecodeError:
        return []


@router.get("/items")
async def query_items(
    limit: int = Query(100, ge=1, le=500),
    offset: int = Query(0, ge=0),
    filters: Optional[str] = Query(None, description='JSON 过滤条件'),
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """分页查询清单项。

    filters 示例：`[["item_name","like","土方"],["data_source_type","=","completed"]]`
    """
    domain = _parse_domain(filters)
    result = search_items(db=db, domain=domain, limit=limit, offset=offset)
    return result


@router.get("/items/{biz_id}")
async def query_item_detail(
    biz_id: str,
    user=Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """按稳定业务 ID 取单条详情。"""
    item = get_item(db, biz_id)
    if item is None:
        raise HTTPException(status_code=404, detail=f"未找到清单项：{biz_id}")
    return item


@router.get("/export")
async def export_excel(
    filters: Optional[str] = Query(None, description='JSON 过滤条件'),
    confirm: bool = Query(False, description='大批量导出二次确认（>1000 行）'),
    user=Depends(require_role(ROLE_ESTIMATOR, ROLE_ADMIN)),
    db: Session = Depends(get_db),
):
    """导出当前筛选结果集为 Excel。

    闸门（M2 §6.2）：
    1. 导出留痕：每次导出写 audit_log（谁、何时、筛选条件、行数、导出字段）；
    2. 大批量闸门：单次 > 1000 行需 confirm=true 二次确认。
    """
    domain = _parse_domain(filters)
    rows = export_items(db=db, domain=domain)

    # 大批量闸门
    if len(rows) > EXPORT_GATE_LIMIT and not confirm:
        raise HTTPException(
            status_code=428,  # Precondition Required
            detail={
                'message': f'本次导出 {len(rows)} 行，超过 {EXPORT_GATE_LIMIT} 行闸门，'
                           '需二次确认（提示：本次导出含商业敏感单价）。',
                'requires_confirm': True,
                'row_count': len(rows),
            },
        )

    # 生成 Excel（内存，不落盘）
    xlsx_bytes = _build_excel(rows)
    filename = f"boq_export_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"

    # 导出留痕（M2 §6.2）
    log_audit(
        db=db,
        model="boq_item",
        res_id=None,
        action=ACTION_EXPORT,
        operator=user.get('username', 'system'),
        reason=f"导出 {len(rows)} 行，筛选条件: {filters or '(全部)'}，字段: {len(EXPORT_COLUMNS)} 列",
        trace_id=uuid.uuid4().hex,
    )
    db.commit()

    return StreamingResponse(
        io.BytesIO(xlsx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


def _build_excel(rows: List[dict]) -> bytes:
    """构建 Excel 字节流（内存）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "清单导出"

    # 表头
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4D7CFE", end_color="4D7CFE", fill_type="solid")
    for col_idx, (_, label) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill

    # 数据行
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, (field, _) in enumerate(EXPORT_COLUMNS, start=1):
            value = row.get(field)
            if value is None:
                value = ''
            ws.cell(row=row_idx, column=col_idx, value=value)

    # 列宽自适应（简单估算）
    for col_idx, (_, label) in enumerate(EXPORT_COLUMNS, start=1):
        max_len = len(label)
        for row_idx in range(2, min(len(rows) + 2, 20) + 1):  # 抽样前 20 行
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    # 允许例外：导出向内存 BytesIO 生成全新导出文件，不触碰导入源文件
    wb.save(buf)
    return buf.getvalue()
