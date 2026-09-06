# -*- coding: utf-8 -*-
"""M2.1 Excel 解析流水线测试（import_service）。

测试用例覆盖：
1. 基本解析（标准表头 + 数据行）
2. 合计行检测（不入库，记录 summary_total）
3. 空行跳过
4. 章节标题行跳过（无编码且无量价）
5. 国标编码解析（parse_gb_code 调用）
6. 单位归一化（normalize_unit 调用）
7. 数值字段解析（含千分位/货币符号容忍）
8. 异常标记（非数值、未归一化单位）
9. 必填列校验（缺少 item_name → ValueError）
10. 未找到表头 → ValueError
11. 多 sheet 指定
12. SHA256 计算稳定
13. preview_parsed 精简信封
"""
import io
import os
import tempfile

import openpyxl
import pytest

from app.services.import_service import (
    parse_excel,
    preview_parsed,
    _to_float,
    _sha256,
)


def _make_excel(rows, sheet_name="Sheet1"):
    """创建内存 Excel → bytes。rows 为 list[list]，第一行是表头。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 标准表头（含全部必填列 + 常用列）
STANDARD_HEADER = [
    '序号', '项目编码', '项目名称', '项目特征描述',
    '计量单位', '工程量', '综合单价', '合价',
]


class TestBasicParse:
    """基本解析测试。"""

    def test_single_data_row(self):
        """单行数据解析：字段映射正确。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '挖一般土方', '土壤类别:二类土', 'm3', 100, 50.5, 5050],
        ]
        result = parse_excel(file_bytes=_make_excel(data), filename="test.xlsx")
        assert result['row_count'] == 1
        row = result['rows'][0]
        assert row['item_name'] == '挖一般土方'
        assert row['item_code'] == '010101001001'  # 规范化后
        assert row['item_code_raw'] == '010101001001'
        assert row['item_feature'] == '土壤类别:二类土'
        assert row['unit'] == 'm3'
        assert row['unit_std'] == 'm³'  # 归一化
        assert row['quantity_num'] == 100.0
        assert row['unit_rate_num'] == 50.5
        assert row['total_num'] == 5050.0
        assert row['anomaly_flag'] == 'normal'

    def test_parsed_total_accumulation(self):
        """多行数据：parsed_total 累加。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '挖一般土方', '', 'm3', 100, 50, 5000],
            [2, '010101002001', '挖沟槽土方', '', 'm3', 50, 60, 3000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        assert result['row_count'] == 2
        assert result['parsed_total'] == 8000.0

    def test_file_hash_stable(self):
        """同一文件内容 → SHA256 稳定。"""
        data = [STANDARD_HEADER, [1, '010101001001', '测试', '', 'm3', 1, 1, 1]]
        content = _make_excel(data)
        r1 = parse_excel(file_bytes=content, filename="a.xlsx")
        r2 = parse_excel(file_bytes=content, filename="b.xlsx")
        assert r1['file_hash'] == r2['file_hash']
        assert len(r1['file_hash']) == 64  # SHA256 十六进制长度


class TestSummaryRow:
    """合计行检测测试。"""

    def test_summary_row_skipped(self):
        """合计行不入库，记录 summary_total。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '挖一般土方', '', 'm3', 100, 50, 5000],
            ['', '', '合计', '', '', '', '', 5000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        assert result['row_count'] == 1  # 合计行不入库
        assert result['summary_total'] == 5000.0

    def test_all_summary_keywords(self):
        """所有合计关键词均被识别。"""
        for kw in ['合计', '总计', '小计', '求和', '总结']:
            data = [
                STANDARD_HEADER,
                [1, '010101001001', '挖一般土方', '', 'm3', 100, 50, 5000],
                ['', '', kw, '', '', '', '', 5000],
            ]
            result = parse_excel(file_bytes=_make_excel(data))
            assert result['row_count'] == 1, f"{kw} 未被识别为合计行"
            assert result['summary_total'] == 5000.0


class TestSkippedRows:
    """跳过行测试。"""

    def test_empty_row_skipped(self):
        """空行（item_name 为空）跳过。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '挖一般土方', '', 'm3', 100, 50, 5000],
            ['', '', '', '', '', '', '', ''],
            [2, '010101002001', '挖沟槽土方', '', 'm3', 50, 60, 3000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        assert result['row_count'] == 2
        assert result['skipped_count'] >= 1

    def test_chapter_title_skipped(self):
        """章节标题行（有名称但无编码且无量价）跳过。"""
        data = [
            STANDARD_HEADER,
            ['', '', '一、土石方工程', '', '', '', '', ''],
            [1, '010101001001', '挖一般土方', '', 'm3', 100, 50, 5000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        assert result['row_count'] == 1
        assert result['rows'][0]['item_name'] == '挖一般土方'


class TestCodeParsing:
    """国标编码解析测试。"""

    def test_code_normalized(self):
        """item_code 经 parse_gb_code 规范化。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001', '挖一般土方', '', 'm3', 100, 50, 5000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        row = result['rows'][0]
        assert row['item_code_raw'] == '010101001'
        # parse_gb_code 返回 (9位, 完整值)，完整值应保留原始或补全
        assert row['item_code'] is not None

    def test_code_with_z_prefix(self):
        """Z 前缀编码（2024 规范）。"""
        data = [
            STANDARD_HEADER,
            [1, 'Z010101001001', '测试项', '', 'm3', 10, 100, 1000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        row = result['rows'][0]
        assert row['item_code_raw'] == 'Z010101001001'
        assert row['item_code'] is not None


class TestUnitNormalization:
    """单位归一化测试。"""

    def test_chinese_unit_normalized(self):
        """中文单位 → 标准单位。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '挖一般土方', '', '立方米', 100, 50, 5000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        row = result['rows'][0]
        assert row['unit'] == '立方米'
        assert row['unit_std'] == 'm³'

    def test_unknown_unit_warned(self):
        """未识别单位 → 原样保留 + anomaly warning。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '测试项', '', '自定义单位', 10, 100, 1000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        row = result['rows'][0]
        assert row['anomaly_flag'] == 'warning'
        assert '计量单位未归一化' in row['anomaly_reason']


class TestNumericParsing:
    """数值字段解析测试。"""

    def test_thousand_separator(self):
        """千分位数字容忍。"""
        assert _to_float('1,234.56') == 1234.56

    def test_currency_symbol(self):
        """货币符号容忍。"""
        assert _to_float('￥5000') == 5000.0
        assert _to_float('¥3000') == 3000.0

    def test_none_and_empty(self):
        """None / 空字符串 → None。"""
        assert _to_float(None) is None
        assert _to_float('') is None
        assert _to_float('   ') is None

    def test_int_float_input(self):
        """int/float 直接转换。"""
        assert _to_float(100) == 100.0
        assert _to_float(50.5) == 50.5

    def test_invalid_string(self):
        """非数值字符串 → None。"""
        assert _to_float('abc') is None

    def test_numeric_fields_in_row(self):
        """行解析中数值字段正确转换。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '测试', '', 'm3', '1,000', '50.5', '50,500'],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        row = result['rows'][0]
        assert row['quantity_num'] == 1000.0
        assert row['unit_rate_num'] == 50.5
        assert row['total_num'] == 50500.0


class TestAnomalyDetection:
    """异常标记测试。"""

    def test_non_numeric_total_warned(self):
        """合价非数值 → anomaly warning。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '测试', '', 'm3', 100, 50, '不是数字'],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        row = result['rows'][0]
        assert row['anomaly_flag'] == 'warning'
        assert 'total 非数值' in row['anomaly_reason']

    def test_normal_row_no_anomaly(self):
        """正常行 anomaly_flag = normal。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '测试', '', 'm3', 100, 50, 5000],
        ]
        result = parse_excel(file_bytes=_make_excel(data))
        assert result['rows'][0]['anomaly_flag'] == 'normal'
        assert result['anomaly_count'] == 0


class TestErrorCases:
    """错误场景测试。"""

    def test_no_file_provided(self):
        """未提供文件 → ValueError。"""
        with pytest.raises(ValueError, match="未提供文件"):
            parse_excel()

    def test_missing_required_column(self):
        """缺少必填列（item_name）→ 表头探测失败（未找到表头行）。

        注：当前 REQUIRED_CANON 只有 item_name，而表头探测已要求 item_name 存在，
        因此 parse_excel 中会先报"未找到表头行"。validate_required 纯函数单独验证。
        """
        from data.aliases import validate_required
        # 纯函数验证：header_map 格式为 {列索引: 规范字段名}
        assert validate_required({0: 'ordinal', 1: 'unit'}) == ['item_name']
        assert validate_required({0: 'item_name'}) == []

        # parse_excel：无 item_name 列 → 未找到表头行
        data = [
            ['序号', '项目编码', '计量单位', '工程量'],
            [1, '010101001001', 'm3', 100],
        ]
        with pytest.raises(ValueError, match="未找到表头行"):
            parse_excel(file_bytes=_make_excel(data))

    def test_no_header_found(self):
        """未找到表头行（无项目名称列）→ ValueError。"""
        data = [
            ['这不是表头', '随便写'],
            [1, 2],
        ]
        with pytest.raises(ValueError, match="未找到表头行"):
            parse_excel(file_bytes=_make_excel(data))

    def test_nonexistent_sheet(self):
        """指定不存在的 sheet → ValueError。"""
        data = [STANDARD_HEADER, [1, '010101001001', '测试', '', 'm3', 1, 1, 1]]
        with pytest.raises(ValueError, match="工作表不存在"):
            parse_excel(file_bytes=_make_excel(data), sheet="不存在的Sheet")


class TestMultiSheet:
    """多 Sheet 测试。"""

    def test_specify_sheet(self):
        """指定 sheet 名解析。"""
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "第一个"
        ws1.append(STANDARD_HEADER)
        ws1.append([1, '010101001001', 'Sheet1数据', '', 'm3', 10, 10, 100])
        ws2 = wb.create_sheet("第二个")
        ws2.append(STANDARD_HEADER)
        ws2.append([1, '010101002001', 'Sheet2数据', '', 'm3', 20, 20, 400])
        buf = io.BytesIO()
        wb.save(buf)
        content = buf.getvalue()

        r1 = parse_excel(file_bytes=content, sheet="第一个")
        assert r1['rows'][0]['item_name'] == 'Sheet1数据'
        assert r1['sheet_name'] == '第一个'

        r2 = parse_excel(file_bytes=content, sheet="第二个")
        assert r2['rows'][0]['item_name'] == 'Sheet2数据'
        assert r2['sheet_name'] == '第二个'

    def test_default_first_sheet(self):
        """不指定 sheet → 默认取第一个。"""
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "第一个"
        ws1.append(STANDARD_HEADER)
        ws1.append([1, '010101001001', '默认Sheet', '', 'm3', 10, 10, 100])
        ws2 = wb.create_sheet("第二个")
        ws2.append(STANDARD_HEADER)
        ws2.append([1, '010101002001', '第二个Sheet', '', 'm3', 20, 20, 400])
        buf = io.BytesIO()
        wb.save(buf)

        result = parse_excel(file_bytes=buf.getvalue())
        assert result['sheet_name'] == '第一个'
        assert result['rows'][0]['item_name'] == '默认Sheet'


class TestPreviewParsed:
    """preview_parsed 精简信封测试。"""

    def test_preview_fields(self):
        """preview_parsed 返回精简字段。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '测试', '', 'm3', 100, 50, 5000],
        ]
        parsed = parse_excel(file_bytes=_make_excel(data))
        preview = preview_parsed(parsed)
        assert 'row_count' in preview
        assert 'skipped_count' in preview
        assert 'anomaly_count' in preview
        assert 'warnings' in preview
        assert 'parsed_total' in preview
        assert 'summary_total' in preview
        assert 'unrecognized_columns' in preview
        # 不包含 rows（精简）
        assert 'rows' not in preview


class TestFilePathInput:
    """file_path 输入测试。"""

    def test_parse_from_file_path(self):
        """从本地文件路径解析。"""
        data = [
            STANDARD_HEADER,
            [1, '010101001001', '文件路径测试', '', 'm3', 100, 50, 5000],
        ]
        content = _make_excel(data)
        tmp_fd, tmp_path = tempfile.mkstemp(suffix='.xlsx')
        os.close(tmp_fd)
        try:
            with open(tmp_path, 'wb') as f:
                f.write(content)
            result = parse_excel(file_path=tmp_path, filename="local.xlsx")
            assert result['row_count'] == 1
            assert result['rows'][0]['item_name'] == '文件路径测试'
        finally:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass  # Windows 文件句柄延迟释放，忽略
