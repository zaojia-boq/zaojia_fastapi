# -*- coding: utf-8 -*-
"""M2-查询导出 API 集成测试（query_api + batch_api）。

测试用例覆盖：
查询：
1. GET /api/query/items —— 分页查询（默认 active=True）
2. GET /api/query/items —— filters 过滤（item_name like）
3. GET /api/query/items/{biz_id} —— 单条详情
4. GET /api/query/export —— 导出 Excel（<1000 行直接导出）
5. GET /api/query/export —— 大批量闸门（>1000 行需 confirm）

批次生命周期：
6. POST soft-delete —— 软删除（批次+清单项进回收站）
7. POST restore —— 还原
8. POST hard-delete —— 硬删除（名称确认 + 快照 + 审计）
9. hard-delete 名称不匹配 → 400
10. hard-delete 宽限期内 → 400
"""
import io
import json
import os
import uuid

import openpyxl
import pytest

from app.config import settings
from app.core.security import get_current_user, ROLE_VIEWER
from app.main import app

DEV_TOKEN = getattr(settings, 'dev_token', 'zaojia-dev-token-2026')
AUTH_HEADER = {"Authorization": f"Bearer {DEV_TOKEN}"}

STANDARD_HEADER = [
    '序号', '项目编码', '项目名称', '项目特征描述',
    '计量单位', '工程量', '综合单价', '合价',
]


def _make_excel_bytes(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _import_batch(client, rows, name="test.xlsx"):
    """上传 + 执行导入，返回批次 ID。"""
    content = _make_excel_bytes(rows)
    upload_resp = client.post(
        "/api/import/upload",
        headers=AUTH_HEADER,
        files={"file": (name, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert upload_resp.status_code == 200
    up = upload_resp.json()
    exec_resp = client.post(
        "/api/import/execute",
        headers=AUTH_HEADER,
        data={"file_token": up['file_token'], "tmp_path": up['tmp_path']},
    )
    assert exec_resp.status_code == 200
    return exec_resp.json()['batch_id']


@pytest.fixture
def sample_rows():
    return [
        STANDARD_HEADER,
        [1, '010101001001', '挖一般土方', '土壤类别:二类土', 'm3', 100, 50.5, 5050],
        [2, '010101002001', '挖沟槽土方', '土壤类别:三类土', 'm3', 50, 60, 3000],
    ]


class TestQueryItems:
    """查询列表测试。"""

    def test_query_empty(self, client):
        """无数据时返回空列表。"""
        resp = client.get("/api/query/items", headers=AUTH_HEADER)
        assert resp.status_code == 200
        data = resp.json()
        assert data['total'] == 0
        assert data['data'] == []

    def test_query_after_import(self, client, sample_rows):
        """导入后查询返回数据。"""
        batch_id = _import_batch(client, sample_rows)
        resp = client.get("/api/query/items", headers=AUTH_HEADER)
        assert resp.status_code == 200
        data = resp.json()
        assert data['total'] == 2
        assert data['data'][0]['item_name'] == '挖一般土方'
        assert data['data'][0]['total_num'] == 5050.0
        assert 'source_location' in data['data'][0]

    def test_query_filter_like(self, client, sample_rows):
        """filters 过滤：item_name like 沟槽。"""
        batch_id = _import_batch(client, sample_rows)
        filters = json.dumps([["item_name", "like", "沟槽"]])
        resp = client.get("/api/query/items", headers=AUTH_HEADER, params={"filters": filters})
        assert resp.status_code == 200
        data = resp.json()
        assert data['total'] == 1
        assert data['data'][0]['item_name'] == '挖沟槽土方'

    def test_query_filter_data_source(self, client, sample_rows):
        """filters 过滤：data_source_type。"""
        batch_id = _import_batch(client, sample_rows)
        filters = json.dumps([["data_source_type", "=", "completed"]])
        resp = client.get("/api/query/items", headers=AUTH_HEADER, params={"filters": filters})
        assert resp.status_code == 200
        assert resp.json()['total'] == 2

    def test_query_requires_auth(self, client):
        """未授权 → 401。"""
        resp = client.get("/api/query/items")
        assert resp.status_code == 401


class TestQueryItemDetail:
    """单条详情测试。"""

    def test_get_item_detail(self, client, sample_rows):
        """按 biz_id 取详情。"""
        batch_id = _import_batch(client, sample_rows)
        # 先查列表拿 biz_id
        resp = client.get("/api/query/items", headers=AUTH_HEADER)
        biz_id = resp.json()['data'][0]['biz_id']

        detail = client.get(f"/api/query/items/{biz_id}", headers=AUTH_HEADER)
        assert detail.status_code == 200
        data = detail.json()
        assert data['item_name'] == '挖一般土方'
        assert data['std_name'] is None
        assert data['active'] is True

    def test_get_nonexistent_item(self, client):
        """不存在的 biz_id → 404。"""
        resp = client.get("/api/query/items/nonexistent-biz-id", headers=AUTH_HEADER)
        assert resp.status_code == 404


class TestQueryExport:
    """导出测试。"""

    def test_export_success(self, client, sample_rows):
        """导出 Excel（<1000 行直接导出）。"""
        batch_id = _import_batch(client, sample_rows)
        resp = client.get("/api/query/export", headers=AUTH_HEADER)
        assert resp.status_code == 200
        assert resp.headers['content-type'].startswith('application/vnd.openxmlformats')
        # 验证内容可解析
        content = resp.content
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        # 表头 + 2 行数据
        assert ws.max_row == 3
        assert ws.cell(row=1, column=1).value == '项目编码'
        assert ws.cell(row=1, column=2).value == '项目名称'
        assert ws.cell(row=1, column=7).value == '合价'
        assert ws.cell(row=2, column=2).value == '挖一般土方'
        assert ws.cell(row=2, column=7).value == 5050 or ws.cell(row=2, column=7).value == '5050'

    def test_export_gate_requires_confirm(self, client):
        """大批量闸门：>1000 行需 confirm=true。"""
        # 构造 1001 行数据
        rows = [STANDARD_HEADER]
        for i in range(1001):
            rows.append([i, f'010101{i:06d}', f'测试项{i}', '', 'm3', 1, 1, 1])
        batch_id = _import_batch(client, rows, name="big.xlsx")

        # 无 confirm → 428
        resp = client.get("/api/query/export", headers=AUTH_HEADER)
        assert resp.status_code == 428
        data = resp.json()['detail']
        assert data['requires_confirm'] is True
        assert data['row_count'] == 1001

        # 带 confirm=true → 200
        resp2 = client.get("/api/query/export", headers=AUTH_HEADER, params={"confirm": "true"})
        assert resp2.status_code == 200

    def test_export_requires_auth(self, client):
        """未授权导出 → 401。"""
        resp = client.get("/api/query/export")
        assert resp.status_code == 401

    def test_export_rejects_viewer(self, client):
        """viewer 角色导出 → 403（M1.5 角色定义：viewer 禁导出）。"""
        # 模拟 viewer 角色（开发期 get_current_user 固定返回 admin，需 override）
        app.dependency_overrides[get_current_user] = lambda: {"username": "test_viewer", "role": ROLE_VIEWER}
        try:
            resp = client.get("/api/query/export", headers=AUTH_HEADER)
            assert resp.status_code == 403
        finally:
            app.dependency_overrides.pop(get_current_user, None)


class TestBatchLifecycle:
    """批次生命周期测试。"""

    def test_soft_delete_and_restore(self, client, sample_rows):
        """软删除 → 清单项消失 → 还原 → 清单项恢复。"""
        batch_id = _import_batch(client, sample_rows)

        # 查询确认有数据
        assert client.get("/api/query/items", headers=AUTH_HEADER).json()['total'] == 2

        # 软删除
        resp = client.post(
            f"/api/import/batches/{batch_id}/soft-delete",
            headers=AUTH_HEADER,
            json={"reason": "测试软删除"},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data['ok'] is True
        assert data['items_soft_deleted'] == 2
        assert 'deleted_at' in data

        # 查询不到（active=False 被默认过滤）
        assert client.get("/api/query/items", headers=AUTH_HEADER).json()['total'] == 0

        # 批次详情显示 inactive
        batch_detail = client.get(f"/api/import/batches/{batch_id}", headers=AUTH_HEADER).json()
        assert batch_detail['active'] is False

        # 还原
        resp2 = client.post(
            f"/api/import/batches/{batch_id}/restore",
            headers=AUTH_HEADER,
        )
        assert resp2.status_code == 200
        assert resp2.json()['items_restored'] == 2

        # 查询恢复
        assert client.get("/api/query/items", headers=AUTH_HEADER).json()['total'] == 2

    def test_soft_delete_twice_rejected(self, client, sample_rows):
        """重复软删除 → 400。"""
        batch_id = _import_batch(client, sample_rows)
        client.post(f"/api/import/batches/{batch_id}/soft-delete", headers=AUTH_HEADER, json={"reason": "t"})
        resp = client.post(f"/api/import/batches/{batch_id}/soft-delete", headers=AUTH_HEADER, json={"reason": "t"})
        assert resp.status_code == 400

    def test_restore_not_in_trash(self, client, sample_rows):
        """未软删的批次还原 → 400。"""
        batch_id = _import_batch(client, sample_rows)
        resp = client.post(f"/api/import/batches/{batch_id}/restore", headers=AUTH_HEADER)
        assert resp.status_code == 400

    def test_hard_delete_wrong_name(self, client, sample_rows):
        """硬删除名称不匹配 → 400。"""
        batch_id = _import_batch(client, sample_rows)
        resp = client.post(
            f"/api/import/batches/{batch_id}/hard-delete",
            headers=AUTH_HEADER,
            json={"confirm_name": "错误名称"},
        )
        assert resp.status_code == 400

    def test_hard_delete_direct_success(self, client, sample_rows):
        """未软删直接硬删除（T9 场景）：名称匹配 + 快照导出 + 物理删除。"""
        batch_id = _import_batch(client, sample_rows)
        # 取批次名
        batch_detail = client.get(f"/api/import/batches/{batch_id}", headers=AUTH_HEADER).json()
        name = batch_detail['name']

        resp = client.post(
            f"/api/import/batches/{batch_id}/hard-delete",
            headers=AUTH_HEADER,
            json={"confirm_name": name},
        )
        assert resp.status_code == 200
        data = resp.json()
        assert data['ok'] is True
        assert data['items_hard_deleted'] == 2
        assert data['snapshot_path'] is not None
        # 快照文件存在
        assert os.path.isfile(data['snapshot_path'])

        # 批次已物理删除 → 404
        assert client.get(f"/api/import/batches/{batch_id}", headers=AUTH_HEADER).status_code == 404
        # 清单项已删除 → 查询为空
        assert client.get("/api/query/items", headers=AUTH_HEADER).json()['total'] == 0

    def test_hard_delete_grace_period(self, client, sample_rows, monkeypatch):
        """宽限期内禁止硬删（软删后 30 天内）。"""
        batch_id = _import_batch(client, sample_rows)
        # 软删
        client.post(f"/api/import/batches/{batch_id}/soft-delete", headers=AUTH_HEADER, json={"reason": "t"})
        batch_detail = client.get(f"/api/import/batches/{batch_id}", headers=AUTH_HEADER).json()
        name = batch_detail['name']

        resp = client.post(
            f"/api/import/batches/{batch_id}/hard-delete",
            headers=AUTH_HEADER,
            json={"confirm_name": name},
        )
        assert resp.status_code == 400
        assert '宽限期' in resp.json()['detail']
